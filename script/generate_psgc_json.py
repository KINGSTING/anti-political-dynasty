import json
import os
import re
from openpyxl import load_workbook

BASE_DIR = r"C:\Users\HP\Desktop\JJ\anti-political-dynasty"
EXCEL_FILE = os.path.join(BASE_DIR, "data", "PSGC-3Q-2026-Publication-Datafile.xlsx")
OUTPUT_JSON = os.path.join(BASE_DIR, "web", "data", "psgc_names.json")

def clean_code(val):
    if val is None:
        return ""
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return re.sub(r'\D', '', s)

def main():
    print(f"Loading {EXCEL_FILE}...")
    wb = load_workbook(EXCEL_FILE, read_only=True, data_only=True)
    ws = wb.active

    code_col = name_col = corr_col = level_col = None
    header_found = False
    psgc_dict = {}
    levels_found = {}
    ilocos_rows = []

    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        # --- Header detection: find row containing "geographic level" ---
        if not header_found:
            row_str = [str(v).strip().lower() if v is not None else "" for v in row]
            if any("geographic level" in v or v.strip() == "level" for v in row_str):
                for j, v in enumerate(row_str):
                    if "10-digit" in v or (v.startswith("psgc") and "standard" not in v):
                        code_col = j
                    if v.strip() == "name":
                        name_col = j
                    if "correspondence" in v:
                        corr_col = j
                    if "geographic level" in v or v.strip() == "level":
                        level_col = j
                if code_col is not None and name_col is not None:
                    header_found = True
                    print(f"✅ Real header at row {row_idx}: code={code_col}, name={name_col}, corr={corr_col}, level={level_col}")
            continue

        # --- Data extraction ---
        code_10 = clean_code(row[code_col]) if code_col is not None and code_col < len(row) else ""
        name = str(row[name_col]).strip() if name_col is not None and name_col < len(row) and row[name_col] is not None else ""
        code_9 = clean_code(row[corr_col]) if corr_col is not None and corr_col < len(row) else ""
        level = str(row[level_col]).strip() if level_col is not None and level_col < len(row) and row[level_col] is not None else ""

        if not code_10 or not name or name.lower() == "nan":
            continue

        levels_found[level] = levels_found.get(level, 0) + 1
        if "ilocos norte" in name.lower():
            ilocos_rows.append((code_10, code_9, name, level))

        # Store 10-digit code
        psgc_dict[code_10] = name

        # Store correspondence code (official 9-digit)
        if code_9 and code_9 != code_10:
            psgc_dict[code_9] = name

        # Also add truncation for non-barangay levels (safety net)
        if level.upper() not in ("BGY", "BARANGAY"):
            truncated = code_10[:-1]
            if truncated not in psgc_dict:
                psgc_dict[truncated] = name

    wb.close()

    if not psgc_dict:
        print("❌ No mappings found.")
        return

    os.makedirs(os.path.dirname(OUTPUT_JSON), exist_ok=True)
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(psgc_dict, f, ensure_ascii=False, indent=2)

    print(f"\n✅ Saved {len(psgc_dict):,} mappings")
    print(f"\n=== Levels found in file ===")
    for lvl, count in sorted(levels_found.items(), key=lambda x: -x[1]):
        print(f"  {lvl or '(blank)'}: {count:,}")

    print(f"\n=== Ilocos Norte rows found ===")
    if ilocos_rows:
        for r in ilocos_rows:
            print(f"  code_10={r[0]} | code_9(corr)={r[1]} | name={r[2]} | level={r[3]}")
    else:
        print("  ⚠️ No rows contain 'Ilocos Norte' in the name column!")

    print(f"\n=== Lookup test ===")
    for test_code in ["012800000", "0128000000", "130000000", "1300000000"]:
        result = psgc_dict.get(test_code, "NOT FOUND")
        print(f"  {test_code}: {result}")

if __name__ == "__main__":
    main()